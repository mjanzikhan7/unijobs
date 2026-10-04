"""Getting plain text out of a PDF or DOCX, and refusing anything that is not one.

Kept apart from :mod:`screening.cv_domain` so that module has zero third-party imports and can be
tested against literal strings. Everything here touches bytes from an untrusted upload, so it is
also where the hostile-input checks live.
"""

from __future__ import annotations

import io
import logging
import zipfile

logger = logging.getLogger(__name__)

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"

MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024

_DOCX_MEMBER = "word/document.xml"


class UnreadableCV(ValueError):
    """The upload is not a readable PDF or DOCX, or is not safe to open."""


def sniff_format(data: bytes) -> str:
    """Return ``"pdf"`` or ``"docx"`` from the bytes themselves.

    Never from the filename or the declared content type: both are attacker-controlled, and
    "trust the extension" is how a renamed payload gets opened by the wrong parser.
    """
    if data.startswith(_PDF_MAGIC):
        return "pdf"
    if data.startswith(_ZIP_MAGIC):
        return "docx"
    raise UnreadableCV("That file is not a PDF or a Word document.")


def _guard_zip_bomb(data: bytes) -> None:
    """Refuse a docx whose contents expand beyond the ceiling."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            total = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as error:
        raise UnreadableCV("That Word document could not be opened.") from error

    if _DOCX_MEMBER not in names:
        raise UnreadableCV("That file is a zip archive, but not a Word document.")

    if total > MAX_UNCOMPRESSED_BYTES:
        logger.warning("rejected a docx expanding to %s bytes", total)
        raise UnreadableCV("That Word document is too large once unpacked.")


def pdf_to_text(data: bytes) -> str:
    """Extract text from a PDF."""
    from pdfminer.high_level import extract_text
    from pdfminer.pdfparser import PDFSyntaxError

    try:
        return extract_text(io.BytesIO(data)) or ""
    except (PDFSyntaxError, ValueError, TypeError, AssertionError) as error:
        raise UnreadableCV("That PDF could not be read.") from error


def docx_to_text(data: bytes) -> str:
    """Extract text from a Word document, including tables.

    Tables matter: plenty of CVs put the entire skills section in one, and reading only
    paragraphs would silently miss it.
    """
    _guard_zip_bomb(data)

    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as error:
        raise UnreadableCV("That Word document could not be read.") from error

    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(part for part in parts if part.strip())


def extract_text(data: bytes) -> str:
    """Return the plain text of an uploaded CV, whatever supported format it is in."""
    return pdf_to_text(data) if sniff_format(data) == "pdf" else docx_to_text(data)
